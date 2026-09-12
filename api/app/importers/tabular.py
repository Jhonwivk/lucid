"""CSV and XLSX parsers with row/column/sheet provenance and visible units."""

from __future__ import annotations

import csv
import importlib.metadata
from io import BytesIO, StringIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from .zip_safety import assert_safe_ooxml

from .common import (
    BLANK_UNKNOWN,
    EVIDENCE_BOUNDARY,
    ImportBudgetError,
    ImportResult,
    ParsedSpan,
    UnreadableMaterialError,
    excel_col,
    is_blank_cell,
    parse_unit_from_header,
    sha256_hex,
    unit_from_value_form,
)

MAX_TABLE_ROWS = 5000
MAX_TABLE_CELLS = 25000
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _parser_version(package: str) -> str:
    try:
        return importlib.metadata.version(package)
    except importlib.metadata.PackageNotFoundError:
        return "unknown"


def _cell_excerpt(header: str, value: str, unit: str | None, *, blank: bool, formula: bool) -> str:
    shown = f"({BLANK_UNKNOWN})" if blank else value
    parts = [f"{header}={shown}" if header else shown]
    if unit:
        parts.append(f"({unit})")
    if formula:
        parts.append("(formula, not recalculated)")
    return " ".join(parts)


def _record_unit(units: list[dict[str, str]], header: str, unit: str, source: str) -> None:
    item = {"header": header, "unit": unit, "source": source}
    if item not in units:
        units.append(item)


def _check_budget(row_count: int, cell_count: int) -> None:
    if row_count > MAX_TABLE_ROWS:
        raise ImportBudgetError(
            f"Table exceeds the {MAX_TABLE_ROWS} row safety limit. "
            "Split the file or reduce the used range; nothing was interpreted as business rules."
        )
    if cell_count > MAX_TABLE_CELLS:
        raise ImportBudgetError(
            f"Table exceeds the {MAX_TABLE_CELLS} cell safety limit. "
            "Split the file or reduce the used range; nothing was interpreted as business rules."
        )


def _decode_csv_text(content: bytes) -> tuple[str, str, str]:
    if b"\x00" in content:
        raise UnreadableMaterialError(
            "CSV contains NUL bytes and cannot be read as text. The file remains unread."
        )
    try:
        return content.decode("utf-8-sig"), "utf-8-sig", "strict"
    except UnicodeDecodeError:
        try:
            return content.decode("utf-8"), "utf-8", "strict"
        except UnicodeDecodeError:
            return content.decode("utf-8", errors="replace"), "utf-8", "replaced"


def parse_csv(filename: str, content: bytes) -> ImportResult:
    text, encoding, decode_status = _decode_csv_text(content)
    sample = text[:4096]
    dialect = csv.excel
    if sample.strip():
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel

    try:
        rows = [list(row) for row in csv.reader(StringIO(text), dialect)]
    except csv.Error as exc:
        raise UnreadableMaterialError(
            f"CSV could not be parsed ({type(exc).__name__}). The file may be malformed."
        ) from exc

    sheet = filename.rsplit(".", 1)[0] or "csv"
    spans: list[ParsedSpan] = []
    headers: list[str] = []
    units: list[dict[str, str]] = []
    blank_cells = 0
    blank_rows = 0
    data_rows = 0
    cell_count = 0

    if rows:
        raw_headers = [str(item) for item in rows[0]]
        headers = raw_headers
        width = max(len(raw_headers), 1)
        _check_budget(len(rows), len(rows) * width)
        for col_idx, raw in enumerate(raw_headers, start=1):
            cell_count += 1
            name, unit = parse_unit_from_header(raw)
            if unit:
                _record_unit(units, raw, unit, "header")
            excerpt = f"header {raw}" + (f" (unit: {unit})" if unit else "")
            spans.append(
                ParsedSpan(
                    locator_kind="cell",
                    sheet=sheet,
                    cell_ref=f"{excel_col(col_idx)}1",
                    excerpt=excerpt,
                )
            )
        last_col = excel_col(width)
        spans.append(
            ParsedSpan(
                locator_kind="cell",
                sheet=sheet,
                cell_ref=f"A1:{last_col}1",
                excerpt=" | ".join(raw_headers) if any(item.strip() for item in raw_headers) else "(blank header row)",
            )
        )

        for row_idx, row in enumerate(rows[1:], start=2):
            data_rows += 1
            values = [str(item) for item in row]
            if len(values) < width:
                values.extend([""] * (width - len(values)))
            width_here = max(width, len(values), 1)
            last = excel_col(width_here)
            if not any(item.strip() for item in values):
                blank_rows += 1
                spans.append(
                    ParsedSpan(
                        locator_kind="cell",
                        sheet=sheet,
                        cell_ref=f"A{row_idx}:{last}{row_idx}",
                        excerpt=f"(blank row; {BLANK_UNKNOWN})",
                    )
                )
                blank_cells += width_here
                cell_count += width_here
                continue
            joined = ",".join(values)
            spans.append(
                ParsedSpan(
                    locator_kind="cell",
                    sheet=sheet,
                    cell_ref=f"A{row_idx}:{last}{row_idx}",
                    excerpt=joined[:500] if joined.strip() else f"(blank row; {BLANK_UNKNOWN})",
                )
            )
            for col_idx in range(1, width_here + 1):
                cell_count += 1
                header = headers[col_idx - 1] if col_idx - 1 < len(headers) else ""
                value = values[col_idx - 1] if col_idx - 1 < len(values) else ""
                blank = is_blank_cell(value)
                if blank:
                    blank_cells += 1
                name, header_unit = parse_unit_from_header(header) if header else ("", None)
                value_unit = unit_from_value_form(value) if not blank else None
                unit = header_unit or value_unit
                if value_unit and not header_unit:
                    _record_unit(units, header or f"{excel_col(col_idx)}", value_unit, "value_form")
                formula = (not blank) and value.lstrip().startswith("=")
                spans.append(
                    ParsedSpan(
                        locator_kind="cell",
                        sheet=sheet,
                        cell_ref=f"{excel_col(col_idx)}{row_idx}",
                        excerpt=_cell_excerpt(header or name, value, unit, blank=blank, formula=formula),
                    )
                )
    _check_budget(max(len(rows), 0), cell_count)

    empty = not rows or (len(rows) == 1 and not any(str(item).strip() for item in rows[0]))
    if empty:
        spans.append(
            ParsedSpan(
                locator_kind="cell",
                sheet=sheet,
                cell_ref="A1",
                excerpt=f"(empty table; {BLANK_UNKNOWN})",
            )
        )

    notes = (
        f"{EVIDENCE_BOUNDARY} "
        "CSV cells keep row/column provenance. "
        "Blank cells stay blank/unknown and are not coerced to 0 or false. "
        "Units are copied only from explicit header or value forms."
    )
    if decode_status != "strict":
        notes += " UTF-8 replacement was used; this is not a pristine decode."

    return ImportResult(
        filename=filename,
        media_type="text/csv",
        kind="table",
        byte_size=len(content),
        checksum=sha256_hex(content),
        notes=notes,
        metadata={
            "parser": "lucid.csv",
            "parser_version": "stdlib.csv",
            "encoding": encoding,
            "decode_status": decode_status,
            "sheet": sheet,
            "sheet_names": [sheet],
            "sheets": [
                {
                    "name": sheet,
                    "hidden": False,
                    "header_row": 1 if rows else None,
                    "row_count": len(rows),
                    "data_row_count": data_rows,
                    "column_count": len(headers),
                    "column_headers": headers,
                    "units": units,
                    "blank_cells": blank_cells,
                    "blank_rows": blank_rows,
                    "formula_cells": 0,
                    "formulas_recalculated": False,
                }
            ],
            "row_count": len(rows),
            "column_count": len(headers),
            "column_headers": headers,
            "units": units,
            "blank_cells": blank_cells,
            "blank_rows": blank_rows,
            "empty": empty,
            "formulas_recalculated": False,
            "semantic_extraction": "not_performed",
            "evidence_role": "evidence_not_instruction",
            "row_budget": MAX_TABLE_ROWS,
            "cell_budget": MAX_TABLE_CELLS,
        },
        spans=spans,
        content=content,
    )


def _xlsx_cell_text(value: object) -> tuple[str, str, bool]:
    if value is None:
        return "", "blank", True
    if isinstance(value, bool):
        return ("true" if value else "false"), "boolean", False
    if isinstance(value, str):
        stripped = value.strip()
        if stripped == "":
            return "", "blank", True
        if stripped.startswith("="):
            return value, "formula", False
        return value, "text", False
    return str(value), "value", False


def parse_xlsx(filename: str, content: bytes) -> ImportResult:
    if not content:
        raise UnreadableMaterialError("XLSX file is empty and cannot be parsed as a workbook.")
    archive = assert_safe_ooxml(content, kind="XLSX")
    archive.close()
    try:
        workbook = load_workbook(BytesIO(content), data_only=False, read_only=False)
    except (InvalidFileException, BadZipFile, KeyError, OSError, ValueError) as exc:
        raise UnreadableMaterialError(
            f"XLSX could not be read ({type(exc).__name__}). "
            "The file may be malformed or is not a supported .xlsx workbook."
        ) from exc

    spans: list[ParsedSpan] = []
    sheets_meta: list[dict[str, object]] = []
    all_units: list[dict[str, str]] = []
    total_cells = 0
    total_rows = 0

    try:
        if not workbook.worksheets:
            raise UnreadableMaterialError("XLSX workbook has no worksheets.")

        for sheet in workbook.worksheets:
            hidden = str(getattr(sheet, "sheet_state", "visible") or "visible") != "visible"
            max_row = int(sheet.max_row or 0)
            max_col = int(sheet.max_column or 0)
            total_rows += max_row
            _check_budget(total_rows, total_cells + max_row * max(max_col, 1))

            headers: list[str] = []
            units: list[dict[str, str]] = []
            blank_cells = 0
            blank_rows = 0
            formula_cells = 0
            hidden_rows: list[int] = []
            hidden_cols: list[str] = []
            data_row_count = 0

            for col_idx in range(1, max(max_col, 0) + 1):
                dimension = sheet.column_dimensions.get(excel_col(col_idx))
                if dimension is not None and getattr(dimension, "hidden", False):
                    hidden_cols.append(excel_col(col_idx))

            if max_row == 0 or max_col == 0:
                spans.append(
                    ParsedSpan(
                        locator_kind="cell",
                        sheet=sheet.title,
                        cell_ref="A1",
                        excerpt=f"(empty sheet; {BLANK_UNKNOWN})",
                    )
                )
                sheets_meta.append(
                    {
                        "name": sheet.title,
                        "hidden": hidden,
                        "header_row": None,
                        "row_count": 0,
                        "data_row_count": 0,
                        "column_count": 0,
                        "column_headers": [],
                        "units": [],
                        "blank_cells": 0,
                        "blank_rows": 0,
                        "formula_cells": 0,
                        "formulas_recalculated": False,
                        "hidden_rows": [],
                        "hidden_columns": hidden_cols,
                        "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
                    }
                )
                continue

            header_cells = next(sheet.iter_rows(min_row=1, max_row=1, max_col=max_col, values_only=False))
            raw_headers: list[str] = []
            for col_idx, cell in enumerate(header_cells, start=1):
                text, _kind, blank = _xlsx_cell_text(cell.value)
                raw_headers.append("" if blank else text)
                name, unit = parse_unit_from_header(text) if not blank else ("", None)
                if unit:
                    _record_unit(units, text, unit, "header")
                    _record_unit(all_units, text, unit, "header")
                excerpt = f"header {text}" + (f" (unit: {unit})" if unit else "")
                if blank:
                    excerpt = f"header ({BLANK_UNKNOWN})"
                spans.append(
                    ParsedSpan(
                        locator_kind="cell",
                        sheet=sheet.title,
                        cell_ref=f"{excel_col(col_idx)}1",
                        excerpt=excerpt,
                    )
                )
            headers = raw_headers
            last_col = excel_col(max(len(raw_headers), 1))
            spans.append(
                ParsedSpan(
                    locator_kind="cell",
                    sheet=sheet.title,
                    cell_ref=f"A1:{last_col}1",
                    excerpt=" | ".join(h for h in raw_headers) if any(raw_headers) else f"(blank header row; {BLANK_UNKNOWN})",
                )
            )
            total_cells += len(raw_headers)

            for row_idx, row in enumerate(
                sheet.iter_rows(min_row=2, max_row=max_row, max_col=max_col, values_only=False),
                start=2,
            ):
                data_row_count += 1
                row_dimension = sheet.row_dimensions.get(row_idx)
                if row_dimension is not None and getattr(row_dimension, "hidden", False):
                    hidden_rows.append(row_idx)
                values: list[str] = []
                kinds: list[str] = []
                blanks: list[bool] = []
                for cell in row:
                    text, kind, blank = _xlsx_cell_text(cell.value)
                    data_type = getattr(cell, "data_type", None)
                    if data_type == "f" or kind == "formula":
                        kind = "formula"
                        if not text.lstrip().startswith("=") and text:
                            text = f"={text}" if not text.startswith("=") else text
                        formula_cells += 1
                    values.append(text)
                    kinds.append(kind)
                    blanks.append(blank)
                    total_cells += 1
                    _check_budget(total_rows, total_cells)
                width = max(len(headers), len(values), 1)
                last = excel_col(width)
                if all(blanks):
                    blank_rows += 1
                    blank_cells += len(values)
                    spans.append(
                        ParsedSpan(
                            locator_kind="cell",
                            sheet=sheet.title,
                            cell_ref=f"A{row_idx}:{last}{row_idx}",
                            excerpt=f"(blank row; {BLANK_UNKNOWN})",
                        )
                    )
                    continue
                spans.append(
                    ParsedSpan(
                        locator_kind="cell",
                        sheet=sheet.title,
                        cell_ref=f"A{row_idx}:{last}{row_idx}",
                        excerpt=",".join(values)[:500],
                    )
                )
                for col_idx in range(1, width + 1):
                    header = headers[col_idx - 1] if col_idx - 1 < len(headers) else ""
                    value = values[col_idx - 1] if col_idx - 1 < len(values) else ""
                    blank = blanks[col_idx - 1] if col_idx - 1 < len(blanks) else is_blank_cell(value)
                    kind = kinds[col_idx - 1] if col_idx - 1 < len(kinds) else "text"
                    if blank:
                        blank_cells += 1
                    name, header_unit = parse_unit_from_header(header) if header else ("", None)
                    value_unit = unit_from_value_form(value) if not blank else None
                    unit = header_unit or value_unit
                    if value_unit and not header_unit:
                        _record_unit(units, header or f"{excel_col(col_idx)}", value_unit, "value_form")
                        _record_unit(all_units, header or f"{excel_col(col_idx)}", value_unit, "value_form")
                    spans.append(
                        ParsedSpan(
                            locator_kind="cell",
                            sheet=sheet.title,
                            cell_ref=f"{excel_col(col_idx)}{row_idx}",
                            excerpt=_cell_excerpt(
                                header or name,
                                value,
                                unit,
                                blank=blank,
                                formula=kind == "formula",
                            ),
                        )
                    )

            sheets_meta.append(
                {
                    "name": sheet.title,
                    "hidden": hidden,
                    "header_row": 1,
                    "row_count": max_row,
                    "data_row_count": data_row_count,
                    "column_count": max_col,
                    "column_headers": headers,
                    "units": units,
                    "blank_cells": blank_cells,
                    "blank_rows": blank_rows,
                    "formula_cells": formula_cells,
                    "formulas_recalculated": False,
                    "hidden_rows": hidden_rows,
                    "hidden_columns": hidden_cols,
                    "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
                }
            )
    finally:
        workbook.close()

    notes = (
        f"{EVIDENCE_BOUNDARY} "
        "XLSX cells keep worksheet name and A1 cell/range provenance. "
        "Formulas are stored as formula text; values were not recalculated. "
        "Blank cells stay blank/unknown and are not coerced to 0 or false. "
        "Hidden sheets/rows/columns are recorded in metadata rather than treated as absent."
    )
    return ImportResult(
        filename=filename,
        media_type=XLSX_MEDIA,
        kind="table",
        byte_size=len(content),
        checksum=sha256_hex(content),
        notes=notes,
        metadata={
            "parser": "lucid.xlsx.openpyxl",
            "parser_version": _parser_version("openpyxl"),
            "sheets": sheets_meta,
            "sheet_names": [item["name"] for item in sheets_meta],
            "hidden_sheet_names": [item["name"] for item in sheets_meta if item.get("hidden")],
            "units": all_units,
            "formulas_recalculated": False,
            "semantic_extraction": "not_performed",
            "evidence_role": "evidence_not_instruction",
            "row_budget": MAX_TABLE_ROWS,
            "cell_budget": MAX_TABLE_CELLS,
            "empty": all(int(item.get("row_count") or 0) == 0 for item in sheets_meta),
        },
        spans=spans,
        content=content,
    )
