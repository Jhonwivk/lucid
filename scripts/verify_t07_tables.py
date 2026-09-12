#!/usr/bin/env python3
"""Focused T07 proof: CSV/XLSX import through the real FastAPI intake.

Uses a temporary SQLite database, a temporary data directory, and a real
FastAPI process. Proves bytes, checksum, sheet+A1 provenance, unit hints,
formula preservation, blank/unknown, duplicate names, and controlled rejects.
"""

from __future__ import annotations

import hashlib
import json
import os
import socket
import sqlite3
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
API_ROOT = ROOT / "api"
VENV_PYTHON = API_ROOT / ".venv" / "bin" / "python"

CSV_BODY = (
    "Room,Capacity (people),Cost [USD],Duration_h,weight_kg,Share %,Notes\n"
    "A,12,40,2,15.5,10%,ok\n"
    "B,,0,3,,,\n"
)


def expect(condition: bool, message: str) -> None:
    if not condition:
        raise SystemExit(f"FAIL: {message}")


def find_port() -> int:
    with socket.socket(AF_INET := __import__("socket").AF_INET, socket.SOCK_STREAM) as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


def json_request(base: str, method: str, path: str, payload: dict | None = None) -> tuple[int, dict | list]:
    headers = {"Accept": "application/json"}
    data = None
    if payload is not None or method in {"POST", "PATCH"}:
        data = json.dumps(payload or {}).encode("utf-8")
        headers["Content-Type"] = "application/json"
    req = urllib.request.Request(base + path, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            raw = resp.read().decode("utf-8")
            return resp.status, json.loads(raw) if raw else {}
    except urllib.error.HTTPError as exc:
        raw = exc.read().decode("utf-8")
        try:
            parsed: dict | list = json.loads(raw) if raw else {}
        except json.JSONDecodeError:
            parsed = {"raw": raw}
        return exc.code, parsed


def import_file(
    base: str,
    project_id: str,
    filename: str,
    content: bytes,
    content_type: str = "application/octet-stream",
) -> tuple[int, dict | list]:
    boundary = "----LucidT07Boundary"
    body = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'
        f"Content-Type: {content_type}\r\n\r\n"
    ).encode("utf-8") + content + f"\r\n--{boundary}--\r\n".encode("utf-8")
    req = urllib.request.Request(
        base + f"/api/projects/{project_id}/materials/import",
        data=body,
        method="POST",
        headers={
            "Accept": "application/json",
            "Content-Type": f"multipart/form-data; boundary={boundary}",
        },
    )
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            raw = resp.read().decode("utf-8")
            return resp.status, json.loads(raw) if raw else {}
    except urllib.error.HTTPError as exc:
        raw = exc.read().decode("utf-8")
        try:
            parsed: dict | list = json.loads(raw) if raw else {}
        except json.JSONDecodeError:
            parsed = {"raw": raw}
        return exc.code, parsed


def wait_health(base: str) -> None:
    deadline = time.time() + 15
    last_error = "timeout"
    while time.time() < deadline:
        try:
            status, body = json_request(base, "GET", "/api/health")
            if status == 200 and isinstance(body, dict) and body.get("status") == "ok":
                return
            last_error = f"status={status} body={body}"
        except Exception as exc:  # noqa: BLE001
            last_error = str(exc)
        time.sleep(0.1)
    raise SystemExit(f"FAIL: API did not become healthy: {last_error}")


def sha256_label(content: bytes) -> str:
    return f"sha256:{hashlib.sha256(content).hexdigest()}"


def _venv_site() -> Path | None:
    lib = API_ROOT / ".venv" / "lib"
    if not lib.is_dir():
        return None
    for child in sorted(lib.iterdir()):
        site = child / "site-packages"
        if site.is_dir():
            return site
    return None


def build_xlsx() -> bytes:
    site = _venv_site()
    if site is not None and str(site) not in sys.path:
        sys.path.insert(0, str(site))
    from openpyxl import Workbook

    workbook = Workbook()
    rooms = workbook.active
    rooms.title = "Rooms"
    rooms.append(["Room", "Capacity (people)", "Cost [USD]", "Duration_h", "weight_kg", "Share %", "Notes"])
    rooms.append(["A", 12, 40, 2, 15.5, "10%", "ok"])
    rooms.append(["B", None, 0, 3, None, None, None])
    rooms["B4"] = "=B2+B3"
    rooms["A4"] = "Total"
    constraints = workbook.create_sheet("Constraints")
    constraints.append(["Rule", "Limit (people)"])
    constraints.append(["No overlap", 1])
    hidden = workbook.create_sheet("HiddenRates")
    hidden.append(["Lane", "Rate [USD]"])
    hidden.append(["night", 80])
    hidden.sheet_state = "hidden"
    buf = BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def list_materials(base: str, project_id: str) -> list[dict]:
    status, body = json_request(base, "GET", f"/api/projects/{project_id}/materials")
    expect(status == 200 and isinstance(body, list), f"list materials {status}")
    return body  # type: ignore[return-value]


def reopen_material(db_path: Path, material_id: str) -> dict:
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        row = conn.execute("SELECT * FROM material WHERE id = ?", (material_id,)).fetchone()
        expect(row is not None, f"material {material_id} in sqlite")
        return dict(row)
    finally:
        conn.close()


def reopen_spans(db_path: Path, material_id: str) -> list[dict]:
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute(
            "SELECT * FROM source_span WHERE material_id = ? ORDER BY created_at, id",
            (material_id,),
        ).fetchall()
        return [dict(row) for row in rows]
    finally:
        conn.close()


def stored_file(data_dir: Path, material: dict) -> Path:
    metadata = material.get("metadata") or {}
    stored = metadata.get("stored_path")
    expect(isinstance(stored, str) and stored, "stored_path in metadata")
    path = data_dir / stored
    expect(path.is_file(), f"stored file exists: {path}")
    return path


def unit_map(metadata: dict) -> dict[str, str]:
    units = metadata.get("units") or []
    mapping: dict[str, str] = {}
    if isinstance(units, dict):
        return {str(k): str(v) for k, v in units.items()}
    for item in units:
        if isinstance(item, dict) and item.get("header") and item.get("unit"):
            mapping[str(item["header"])] = str(item["unit"])
    return mapping


def main() -> None:
    tmp = tempfile.TemporaryDirectory(prefix="lucid-t07-")
    tmp_path = Path(tmp.name)
    db_path = tmp_path / "lucid.db"
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    port = find_port()
    base = f"http://127.0.0.1:{port}"
    env = os.environ.copy()
    env["LUCID_DB_PATH"] = str(db_path)
    env["LUCID_DATA_DIR"] = str(data_dir)
    python = str(VENV_PYTHON if VENV_PYTHON.exists() else sys.executable)
    proc = subprocess.Popen(
        [
            python,
            "-m",
            "uvicorn",
            "app.main:app",
            "--app-dir",
            str(API_ROOT),
            "--host",
            "127.0.0.1",
            "--port",
            str(port),
            "--log-level",
            "warning",
        ],
        cwd=str(API_ROOT),
        env=env,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    try:
        wait_health(base)
        status, shell = json_request(base, "GET", "/api/shell")
        expect(status == 200 and isinstance(shell, dict), "shell")
        expect(shell["capabilities"]["csv_xlsx"] == "implemented", "csv_xlsx capability")
        expect("csv" in shell["capabilities"]["import"], "import capability lists csv")

        status, project = json_request(
            base,
            "POST",
            "/api/projects",
            {"title": "T07 table intake proof", "summary": "csv/xlsx provenance"},
        )
        expect(status == 200 and isinstance(project, dict), f"create project {status} {project}")
        project_id = project["id"]

        csv_bytes = CSV_BODY.encode("utf-8")
        status, csv_import = import_file(base, project_id, "rooms.csv", csv_bytes, "text/csv")
        expect(status == 200 and isinstance(csv_import, dict), f"csv import {status} {csv_import}")
        csv_material = csv_import["material"]
        csv_spans = csv_import["spans"]
        expect(csv_material["filename"] == "rooms.csv", "csv original filename")
        expect(csv_material["kind"] == "table", "csv kind")
        expect(csv_material["byte_size"] == len(csv_bytes), "csv byte_size")
        expect(csv_material["checksum"] == sha256_label(csv_bytes), "csv sha256")
        expect(stored_file(data_dir, csv_material).read_bytes() == csv_bytes, "stored csv bytes")
        db_row = reopen_material(db_path, csv_material["id"])
        expect(db_row["checksum"] == sha256_label(csv_bytes), "sqlite csv checksum")
        db_spans = reopen_spans(db_path, csv_material["id"])
        expect(len(db_spans) == len(csv_spans), "sqlite span count matches")
        cell_spans = [span for span in csv_spans if span.get("locator_kind") == "cell"]
        expect(any(span.get("cell_ref") == "B2" and span.get("sheet") == "rooms" for span in cell_spans), "csv B2")
        b3 = next(span for span in cell_spans if span.get("cell_ref") == "B3")
        expect("blank/unknown" in (b3.get("excerpt") or ""), f"blank capacity stays unknown: {b3.get('excerpt')}")
        expect("=0" not in (b3.get("excerpt") or "") and "B3=0" not in (b3.get("excerpt") or ""), "blank not coerced to 0")
        c3 = next(span for span in cell_spans if span.get("cell_ref") == "C3")
        expect("=0" in (c3.get("excerpt") or "") or "0" in (c3.get("excerpt") or ""), f"known zero preserved: {c3}")
        units = unit_map(csv_material["metadata"])
        expect(units.get("Capacity (people)") == "people", f"unit people {units}")
        expect(units.get("Cost [USD]") == "USD", f"unit USD {units}")
        expect(units.get("Duration_h") == "h", f"unit h {units}")
        expect(units.get("weight_kg") == "kg", f"unit kg {units}")
        expect(units.get("Share %") == "%", f"unit percent {units}")

        xlsx_bytes = build_xlsx()
        status, xlsx_import = import_file(
            base,
            project_id,
            "rooms.xlsx",
            xlsx_bytes,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        expect(status == 200 and isinstance(xlsx_import, dict), f"xlsx import {status} {xlsx_import}")
        xlsx_material = xlsx_import["material"]
        xlsx_meta = xlsx_material["metadata"]
        expect(xlsx_material["kind"] == "table", "xlsx kind")
        expect(xlsx_material["checksum"] == sha256_label(xlsx_bytes), "xlsx sha256")
        expect(stored_file(data_dir, xlsx_material).read_bytes() == xlsx_bytes, "stored xlsx bytes")
        expect(xlsx_meta.get("formulas_recalculated") is False, "formulas not recalculated")
        names = xlsx_meta.get("sheet_names") or []
        expect("Rooms" in names and "Constraints" in names, f"two sheets {names}")
        expect("HiddenRates" in names, f"hidden sheet recorded {names}")
        expect("HiddenRates" in (xlsx_meta.get("hidden_sheet_names") or []), "hidden flag")
        xlsx_spans = xlsx_import["spans"]
        formula = next(
            (
                span
                for span in xlsx_spans
                if span.get("sheet") == "Rooms" and span.get("cell_ref") == "B4"
            ),
            None,
        )
        expect(formula is not None, "Rooms!B4 span")
        excerpt = formula.get("excerpt") or ""
        expect("=B2+B3" in excerpt, f"formula text preserved: {excerpt}")
        expect("not recalculated" in excerpt, f"formula honesty: {excerpt}")
        blank = next(span for span in xlsx_spans if span.get("sheet") == "Rooms" and span.get("cell_ref") == "B3")
        expect("blank/unknown" in (blank.get("excerpt") or ""), f"xlsx blank: {blank.get('excerpt')}")
        zero = next(span for span in xlsx_spans if span.get("sheet") == "Rooms" and span.get("cell_ref") == "C3")
        expect("0" in (zero.get("excerpt") or ""), f"xlsx known zero: {zero.get('excerpt')}")
        constraint = next(
            span for span in xlsx_spans if span.get("sheet") == "Constraints" and span.get("cell_ref") == "B2"
        )
        expect(constraint.get("cell_ref") == "B2", "second sheet A1 provenance")
        xlsx_units = unit_map(xlsx_meta)
        expect(xlsx_units.get("Capacity (people)") == "people", xlsx_units)
        expect(xlsx_units.get("Cost [USD]") == "USD", xlsx_units)
        expect(xlsx_units.get("Duration_h") == "h", xlsx_units)
        expect(xlsx_units.get("weight_kg") == "kg", xlsx_units)
        expect(xlsx_units.get("Share %") == "%", xlsx_units)
        db_xlsx_spans = reopen_spans(db_path, xlsx_material["id"])
        expect(any(row["sheet"] == "Rooms" and row["cell_ref"] == "B4" for row in db_xlsx_spans), "sqlite formula span")

        first_bytes = b"Room,Capacity (people)\nA,1\n"
        second_bytes = b"Room,Capacity (people)\nB,9\n"
        status, first_import = import_file(base, project_id, "rooms.csv", first_bytes, "text/csv")
        expect(status == 200, f"first duplicate csv {status}")
        status, second_import = import_file(base, project_id, "rooms.csv", second_bytes, "text/csv")
        expect(status == 200, f"second duplicate csv {status}")
        first_mat = first_import["material"]
        second_mat = second_import["material"]
        expect(first_mat["id"] != second_mat["id"], "duplicate name creates two rows")
        expect(first_mat["checksum"] == sha256_label(first_bytes), "first checksum")
        expect(second_mat["checksum"] == sha256_label(second_bytes), "second checksum")
        first_path = stored_file(data_dir, first_mat)
        second_path = stored_file(data_dir, second_mat)
        expect(first_path != second_path, "collision-safe distinct stored paths")
        expect(first_path.read_bytes() == first_bytes, "first stored bytes intact")

        before_reject = len(list_materials(base, project_id))
        status, xls_body = import_file(base, project_id, "legacy.xls", b"not-xlsx", "application/vnd.ms-excel")
        expect(status == 415, f"xls 415, got {status} {xls_body}")
        status, bad_xlsx = import_file(
            base,
            project_id,
            "broken.xlsx",
            b"PK\x03\x04not-a-workbook",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        expect(status == 422, f"malformed xlsx 422, got {status} {bad_xlsx}")
        oversized = b"a,b\n" + (b"x," * 1000) + b"x\n"  # not oversize; real oversize:
        huge = b"a,b\n" + (b"1,2\n" * 10)
        huge = b"x" * (12 * 1024 * 1024 + 1)
        status, big_body = import_file(base, project_id, "huge.csv", huge, "text/csv")
        expect(status == 413, f"oversize 413, got {status} {big_body}")
        expect(len(list_materials(base, project_id)) == before_reject, "rejects did not persist")

        print("PASS: T07 CSV/XLSX import through real FastAPI + SQLite")
        print(
            json.dumps(
                {
                    "project_id": project_id,
                    "csv_material": csv_material["id"],
                    "xlsx_material": xlsx_material["id"],
                    "xlsx_sheets": names,
                    "materials": len(list_materials(base, project_id)),
                },
                sort_keys=True,
            )
        )
        print(f"temp_db={db_path}")
        print(f"temp_data={data_dir}")
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=5)
        except subprocess.TimeoutExpired:
            proc.kill()
        tmp.cleanup()


if __name__ == "__main__":
    main()
